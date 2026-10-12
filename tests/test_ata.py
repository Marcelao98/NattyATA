import tempfile
import unittest
from contextlib import redirect_stderr, redirect_stdout
from datetime import date
from pathlib import Path
import io
import json

from vatty import ata, cli

try:
    import docx
except ImportError:  # pragma: no cover
    docx = None
try:
    import reportlab  # noqa: F401
except ImportError:  # pragma: no cover
    reportlab = None

MD = """# Título que o código ignora

## Resumo
- Fechar a proposta. **Prazo:** sexta.

## Decisões
- [00:32] Fechar na sexta (Pessoa 3)
  - Depende do orçamento

## Tarefas
- Pessoa 2 | Revisar o orçamento | quinta | [00:07]
- Pessoa 10 | Enviar a proposta

## Pendências
- [00:24] Custo: falta aprovar
"""


class TestLer(unittest.TestCase):
    def test_secoes_itens_e_subitens(self):
        secoes = ata.ler(MD)
        self.assertEqual([s.titulo for s in secoes], ["Resumo", "Decisões", "Tarefas", "Pendências"])
        self.assertEqual(secoes[1].itens[0].filhos, ["Depende do orçamento"])
        self.assertEqual(len(secoes[2].itens), 2)

    def test_ignora_titulo_e_linhas_vazias_e_descarta_secao_sem_itens(self):
        secoes = ata.ler("# T\n\n## Vazia\n\n## Cheia\n- a\n")
        self.assertEqual([s.titulo for s in secoes], ["Cheia"])

    def test_texto_solto_antes_de_qualquer_secao_vira_resumo(self):
        self.assertEqual(ata.ler("só uma linha\n")[0].titulo, "Resumo")

    def test_vazio(self):
        self.assertEqual(ata.ler(""), [])


class TestNomes(unittest.TestCase):
    def test_troca_palavra_inteira(self):
        nomes = {"Pessoa 1": "Ana", "Pessoa 2": "Beto"}
        self.assertEqual(ata.aplicar_nomes("Pessoa 1, Pessoa 10 e Pessoa 2.", nomes), "Ana, Pessoa 10 e Beto.")

    def test_ler_nomes(self):
        self.assertEqual(ata.ler_nomes(["Pessoa 1=Ana", " Pessoa 2 = Beto "]), {"Pessoa 1": "Ana", "Pessoa 2": "Beto"})
        self.assertEqual(ata.ler_nomes(None), {})
        for ruim in ("Pessoa 1", "=Ana", "Pessoa 1="):
            with self.assertRaises(ata.AtaErro):
                ata.ler_nomes([ruim])


class TestTabela(unittest.TestCase):
    def test_tarefas_viram_linhas_de_quatro_campos(self):
        tarefas = ata.ler(MD)[2]
        self.assertTrue(ata.eh_tabela(tarefas))
        self.assertEqual(ata.linhas_da_tabela(tarefas),
                         [("Pessoa 2", "Revisar o orçamento", "quinta", "[00:07]"),
                          ("Pessoa 10", "Enviar a proposta", "—", "—")])

    def test_so_a_secao_tarefas_e_tabela(self):
        self.assertFalse(ata.eh_tabela(ata.Secao("Decisões", [ata.Item("a | b")])))
        self.assertFalse(ata.eh_tabela(ata.Secao("Tarefas", [ata.Item("sem barra")])))


class TestCabecalho(unittest.TestCase):
    def test_gravacao_com_falantes(self):
        fonte = {"titulo": "Gravação x", "data": "2026-10-12_14-30", "duracao_s": 3900}
        falantes = {"pessoas": [{"pessoa": "Pessoa 1", "fala_s": 600.0}, {"pessoa": "Pessoa 2", "fala_s": 10.0}]}
        cab = ata.cabecalho(fonte, falantes, nomes={"Pessoa 1": "Ana"})
        self.assertEqual(cab["titulo"], "Ata: Gravação x")
        d = dict(cab["linhas"])
        self.assertEqual(d["Data"], "12/10/2026 14:30")
        self.assertEqual(d["Duração"], "1 h 05 min")
        self.assertEqual(d["Participantes (pela voz)"], "Ana (10 min de fala); Pessoa 2 (fala curta)")

    def test_sem_falantes_nem_faixas_nao_inventa_participantes(self):
        cab = ata.cabecalho({"titulo": "Aula"}, None, titulo="Minha ata", hoje=date(2026, 10, 12))
        self.assertEqual(cab["titulo"], "Minha ata")
        self.assertEqual(cab["linhas"], [("Data", "12/10/2026")])

    def test_duas_faixas_sem_separacao(self):
        cab = ata.cabecalho({"faixas": ["sistema", "microfone"]}, None, hoje=date(2026, 1, 2))
        self.assertIn("Eu (microfone)", dict(cab["linhas"])["Participantes (pela voz)"])

    def test_data_legivel(self):
        self.assertEqual(ata._data_legivel("2026-10-12_14-30-2"), "12/10/2026 14:30")
        self.assertEqual(ata._data_legivel("ontem"), "ontem")


class TestGerar(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.pasta = Path(tmp.name)
        (self.pasta / "ata-proposta.md").write_text(MD, encoding="utf-8")
        (self.pasta / "fonte.json").write_text(json.dumps(
            {"tipo": "gravacao", "titulo": "Reunião", "duracao_s": 600}), encoding="utf-8")
        (self.pasta / "falantes.json").write_text(json.dumps(
            {"pessoas": [{"pessoa": "Pessoa 2", "fala_s": 300.0}], "turnos": []}), encoding="utf-8")

    @unittest.skipIf(docx is None, "python-docx não instalado")
    def test_docx_tem_cabecalho_tabela_e_nomes(self):
        (arquivo,) = ata.gerar(self.pasta, nomes={"Pessoa 2": "Beto", "Pessoa 3": "Carla"})
        self.assertEqual(arquivo.name, "ata-proposta.docx")
        d = docx.Document(str(arquivo))
        textos = [p.text for p in d.paragraphs]
        self.assertEqual(textos[0], "Ata: Reunião")
        self.assertIn("Duração: 10 min", textos)
        self.assertTrue(any("Beto (5 min de fala)" in t for t in textos))
        self.assertIn("[00:32] Fechar na sexta (Carla)", textos)
        self.assertIn("Prazo: sexta.", " ".join(textos))  # o negrito vira formatação, não asteriscos
        self.assertNotIn("**", " ".join(textos))
        linhas = [[c.text for c in r.cells] for r in d.tables[0].rows]
        self.assertEqual(linhas[0], ["Responsável", "Tarefa", "Prazo", "Horário"])
        self.assertEqual(linhas[1][0], "Beto")
        self.assertEqual(linhas[2][0], "Pessoa 10")  # só palavra inteira
        self.assertNotIn("Título que o código ignora", textos)

    @unittest.skipIf(reportlab is None, "reportlab não instalado")
    def test_pdf_e_gerado(self):
        (arquivo,) = ata.gerar(self.pasta, formatos=("pdf",))
        self.assertTrue(arquivo.read_bytes().startswith(b"%PDF"))

    def test_o_md_do_agente_nao_e_alterado(self):
        antes = (self.pasta / "ata-proposta.md").read_text(encoding="utf-8")
        if docx is not None:
            ata.gerar(self.pasta, nomes={"Pessoa 2": "Beto"})
        self.assertEqual((self.pasta / "ata-proposta.md").read_text(encoding="utf-8"), antes)

    def test_usa_a_ata_mais_recente_ou_a_pedida(self):
        import os
        nova = self.pasta / "ata-outra.md"
        nova.write_text("## Resumo\n- x\n", encoding="utf-8")
        os.utime(self.pasta / "ata-proposta.md", (1, 1))
        self.assertEqual(ata.achar_ata(self.pasta).name, "ata-outra.md")
        self.assertEqual(ata.achar_ata(self.pasta, "ata-proposta.md").name, "ata-proposta.md")

    def test_erros_claros(self):
        with tempfile.TemporaryDirectory() as vazia:
            with self.assertRaises(ata.AtaErro) as ctx:
                ata.gerar(vazia)
            self.assertIn("Secretário", str(ctx.exception))
        with self.assertRaises(ata.AtaErro):
            ata.gerar(self.pasta, "nao-existe.md")
        with self.assertRaises(ata.AtaErro):
            ata.gerar(self.pasta, formatos=("odt",))
        (self.pasta / "ata-ruim.md").write_text("\n\n", encoding="utf-8")
        with self.assertRaises(ata.AtaErro) as ctx:
            ata.gerar(self.pasta, "ata-ruim.md")
        self.assertIn("fora do formato", str(ctx.exception))

    def test_sem_biblioteca_o_erro_diz_como_instalar(self):
        import sys
        from unittest import mock
        with mock.patch.dict(sys.modules, {"docx": None, "reportlab": None, "reportlab.lib": None}):
            for formato in ("docx", "pdf"):
                with self.assertRaises(ata.AtaErro) as ctx:
                    ata.gerar(self.pasta, formatos=(formato,))
                self.assertIn('".[ata]"', str(ctx.exception))


class TestCli(unittest.TestCase):
    def test_comando_ata(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / "ata-x.md").write_text(MD, encoding="utf-8")
            saida, erro = io.StringIO(), io.StringIO()
            with redirect_stdout(saida), redirect_stderr(erro):
                codigo = cli.main(["ata", tmp, "--formato", "docx", "--nomes", "Pessoa 2=Beto"])
            if docx is None:
                self.assertEqual(codigo, 1)
            else:
                self.assertEqual(codigo, 0, erro.getvalue())
                self.assertIn("Ata pronta", saida.getvalue())
                self.assertTrue((Path(tmp) / "ata-x.docx").is_file())

    def test_erro_vira_mensagem_e_codigo_1(self):
        with tempfile.TemporaryDirectory() as tmp:
            erro = io.StringIO()
            with redirect_stderr(erro):
                self.assertEqual(cli.main(["ata", tmp]), 1)
            self.assertIn("Secretário", erro.getvalue())
            with redirect_stderr(erro):
                self.assertEqual(cli.main(["ata", tmp, "--nomes", "sem-igual"]), 1)


if __name__ == "__main__":
    unittest.main()
